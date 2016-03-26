#!/usr/bin/python
#coding:utf-8

import docx
import xlrd
import xlwt
import os
from xlutils.copy import copy as xlcopy

def read_all_docxs(path):
    os.chdir(path)
    files=[]
    for dirpath,dirnames,filenames in os.walk(path):
        files+=filenames

    return files

def read_one_docx(file):
    keys=['company_name','job_in_company','company_industry','company_setup_time','company_major_job','company_business_model','company_cash','company_for_who','company_development_stage','company_incash_way','company_incash_number','company_city','company_location','company_empolyees_number']

    os.chdir('/var/ftpadmin/forwenjie/file_input')
    document=docx.Document(file)
    table_company=document.tables[1]
    #forinsert
    keys_loc_in_table=[(0,1),(99,99),(2,3),(0,3),(2,1),(3,1),(1,3),(99,99),(99,99),(99,99),(99,99),(1,1),(1,1),(3,3)]
    company_info_dict={}

    for i,tup in enumerate(keys_loc_in_table):
        if (tup[0]==99 and tup[1]==99):
            company_info_dict[keys[i]]=None
        else:
            company_info_dict[keys[i]]=table_company.cell(tup[0],tup[1]).text

    return company_info_dict

def main():

    keys=['company_name','job_in_company','company_industry','company_setup_time','company_major_job','company_business_model','company_cash','company_for_who','company_development_stage','company_incash_way','company_incash_number','company_city','company_location','company_empolyees_number']
    path_input=r'/var/ftpadmin/forwenjie/file_input'
    path_output=r'/var/ftpadmin/forwenjie/file_output'

    files=read_all_docxs(path_input)
    path_saving_xls=path_output+'/output.xls'

    read_book=xlrd.open_workbook(path_saving_xls,formatting_info=True)
    write_book=xlcopy(read_book)
    sheet1=write_book.get_sheet(0)
    insert_index=range(9,23)
    for i,file in enumerate(files):
        company_info_dict={}
        company_info_dict=read_one_docx(file)
        for j,key in enumerate(keys):
            sheet1.write(i+1,insert_index[j],company_info_dict[key])
    write_book.save(path_saving_xls)


if __name__=='__main__':
    main()
